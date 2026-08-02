#!/usr/bin/env python3
"""
WEBCLAIR — Fusion du lot (WEBCLAIR_modele.xlsx) dans le registre maître
=========================================================================

But : une fois que l'historien (Claude) a rempli WEBCLAIR_modele.xlsx pour
le lot du jour, ce script recopie les colonnes remplies (Année, Statut,
Descriptions, Sources, Relations, Notes) dans le fichier maître
WEBCLAIR_definitions.xlsx, en s'appuyant sur l'ID.

Sécurité : ne touche QUE les lignes du maître actuellement au statut
"⬜ À faire". Les lignes déjà "✅ Fait" dans le maître ne sont jamais
écrasées, même si elles apparaissent aussi (en exemples) dans le modèle.

Usage :
    python3 merge_modele_into_master.py
"""

import openpyxl
from pathlib import Path

MASTER_PATH = "WEBCLAIR_definitions.xlsx"
MODELE_PATH = "WEBCLAIR_modele.xlsx"
SHEET_MASTER = "Définitions WEBCLAIR"
SHEET_MODELE = "Feuil1"

COLS_TO_COPY = [3, 5, 6, 7, 8, 9, 10]  # Année, Statut, Desc courte, Desc longue, Sources, Relations, Notes
# (0=ID, 1=Terme, 2=Catégorie sont considérés stables, non recopiés)


def main():
    for p in (MASTER_PATH, MODELE_PATH):
        if not Path(p).exists():
            raise FileNotFoundError(f"{p} introuvable dans le dossier courant.")

    wb_modele = openpyxl.load_workbook(MODELE_PATH, data_only=True)
    ws_modele = wb_modele[SHEET_MODELE]

    modele_by_id = {}
    for row in ws_modele.iter_rows(min_row=2, values_only=True):
        rid = row[0]
        if rid:
            modele_by_id[rid] = row

    wb_master = openpyxl.load_workbook(MASTER_PATH)
    ws_master = wb_master[SHEET_MASTER]

    updated, still_pending, not_found_in_modele = 0, 0, 0

    for row in ws_master.iter_rows(min_row=2):
        rid = row[0].value
        statut = row[5].value
        if statut != "⬜ À faire":
            continue  # on ne touche jamais une ligne déjà Fait (ou X, etc.)

        if rid not in modele_by_id:
            not_found_in_modele += 1
            continue

        src = modele_by_id[rid]
        if src[5] != "✅ Fait":
            still_pending += 1
            continue  # pas encore rempli côté modèle, on laisse tel quel

        for col_idx in COLS_TO_COPY:
            row[col_idx].value = src[col_idx]
        updated += 1

    wb_master.save(MASTER_PATH)

    print("=== Fusion modèle -> maître ===")
    print(f"  Lignes mises à jour           : {updated}")
    print(f"  Encore 'À faire' côté modèle  : {still_pending}")
    print(f"  Absentes du modèle (ignorées) : {not_found_in_modele}")
    print(f"✅ {MASTER_PATH} mis à jour.")


if __name__ == "__main__":
    main()
